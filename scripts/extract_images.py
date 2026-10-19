import os
import pptx
from pptx import Presentation

pptx_path = r"C:\tmp\00-Mycat\Engineering\AI in Engineering Design.pptx"
prs = Presentation(pptx_path)
out_dir = r"C:\tmp\00-Mycat\Engineering\assets"
os.makedirs(out_dir, exist_ok=True)

img_count = 0
for i, slide in enumerate(prs.slides):
    for j, shape in enumerate(slide.shapes):
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            ext = img.ext
            filename = f"slide_{i+1}_img_{j+1}.{ext}"
            filepath = os.path.join(out_dir, filename)
            with open(filepath, "wb") as f:
                f.write(img.blob)
            img_count += 1
            print(f"Saved: {filename} ({len(img.blob)} bytes) from Slide {i+1}")

print(f"Total extracted images: {img_count}")
