import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def embed_images_in_pptx():
    pptx_path = r"C:\tmp\00-Mycat\Engineering\Automotive_Engineering_Master_Presentation_32Slides.pptx"
    assets_dir = r"C:\tmp\00-Mycat\Engineering\assets"

    prs = Presentation(pptx_path)
    print("Loaded presentation with", len(prs.slides), "slides.")

    # Image 1: Slide 6 (CAD Evolution)
    # Slide index 5
    s6 = prs.slides[5]
    img6 = os.path.join(assets_dir, "cad_design_evolution.jpg")
    if os.path.exists(img6):
        # We place it in the lower half of the slide
        s6.shapes.add_picture(img6, Inches(0.8), Inches(4.5), width=Inches(11.733), height=Inches(2.55))
        print("Embedded cad_design_evolution.jpg into Slide 6")

    # Image 2: Slide 16 (Flowchart)
    # Slide index 15
    s16 = prs.slides[15]
    img16 = os.path.join(assets_dir, "automotive_flowchart_diagram.jpg")
    if os.path.exists(img16):
        # Let's add a dedicated high-impact visual slide or replace/augment
        # Let's place the image in Slide 16 prominently
        s16.shapes.add_picture(img16, Inches(1.2), Inches(1.6), width=Inches(10.933), height=Inches(5.4))
        print("Embedded automotive_flowchart_diagram.jpg into Slide 16")

    # Image 3: Slide 22 (GD&T Metrology)
    # Slide index 21
    s22 = prs.slides[21]
    img22 = os.path.join(assets_dir, "dimensional_gdt_metrology.jpg")
    if os.path.exists(img22):
        # Place in right column or lower card
        # Let's add it neatly
        s22.shapes.add_picture(img22, Inches(7.8), Inches(1.75), width=Inches(4.733), height=Inches(5.2))
        print("Embedded dimensional_gdt_metrology.jpg into Slide 22")

    # Image 4: Slide 25 (Crash Safety & EV Battery)
    # Slide index 24
    s25 = prs.slides[24]
    img25 = os.path.join(assets_dir, "biw_crash_safety.jpg")
    if os.path.exists(img25):
        s25.shapes.add_picture(img25, Inches(7.8), Inches(1.75), width=Inches(4.733), height=Inches(5.2))
        print("Embedded biw_crash_safety.jpg into Slide 25")

    new_out = r"C:\tmp\00-Mycat\Engineering\Automotive_Engineering_32Slides_Visual.pptx"
    prs.save(new_out)
    print("Successfully saved updated PPTX with embedded high-res illustrations to:", new_out)

if __name__ == "__main__":
    embed_images_in_pptx()
